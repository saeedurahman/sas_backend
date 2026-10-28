"""Items 10 and 11: middleware, logging redaction, CSV/Excel injection (no DB needed)."""

from __future__ import annotations

import csv
import io
import logging

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.logging_config import JsonFormatter, redact
from app.core.middleware import (
    BodySizeLimitMiddleware,
    RequestContextMiddleware,
    SecurityHeadersMiddleware,
)
from app.services.export_service import sanitize_cell


def _app(max_bytes=100, hsts=False):
    app = FastAPI()

    @app.post("/echo")
    async def echo(request: Request):
        body = await request.body()
        return {"n": len(body)}

    @app.get("/ping")
    def ping():
        return {"ok": True}

    app.add_middleware(BodySizeLimitMiddleware, max_bytes=max_bytes)
    app.add_middleware(SecurityHeadersMiddleware, hsts=hsts)
    app.add_middleware(RequestContextMiddleware)
    return app


def test_security_headers_present():
    r = TestClient(_app()).get("/ping")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["x-frame-options"] == "DENY"
    assert r.headers["referrer-policy"] == "no-referrer"
    assert "frame-ancestors 'none'" in r.headers["content-security-policy"]
    assert r.headers["cache-control"] == "no-store"
    assert "strict-transport-security" not in r.headers


def test_hsts_only_when_enabled():
    r = TestClient(_app(hsts=True)).get("/ping")
    assert "max-age=31536000" in r.headers["strict-transport-security"]


def test_request_id_generated_and_echoed():
    client = TestClient(_app())
    generated = client.get("/ping").headers["x-request-id"]
    assert len(generated) == 32
    echoed = client.get("/ping", headers={"X-Request-ID": "trace-12345678"})
    assert echoed.headers["x-request-id"] == "trace-12345678"
    hostile = client.get("/ping", headers={"X-Request-ID": "bad id\twith spaces"})
    assert hostile.headers["x-request-id"] != "bad id\twith spaces"


def test_body_within_limit_ok():
    assert TestClient(_app()).post("/echo", content=b"x" * 100).json() == {"n": 100}


def test_declared_oversize_body_rejected_with_413():
    r = TestClient(_app()).post("/echo", content=b"x" * 101)
    assert r.status_code == 413
    assert r.json() == {"detail": "Request body too large"}


def test_streamed_oversize_body_rejected_with_413():
    def chunks():
        for _ in range(10):
            yield b"x" * 50  # no Content-Length: chunked

    r = TestClient(_app()).post("/echo", content=chunks())
    assert r.status_code == 413


def test_access_log_has_path_but_no_query_or_auth(caplog):
    with caplog.at_level(logging.INFO, logger="app.access"):
        TestClient(_app()).get(
            "/ping?token=SECRET123&pin=999999", headers={"Authorization": "Bearer abc.def.ghi"}
        )
    rec = next(r for r in caplog.records if r.name == "app.access")
    assert rec.path == "/ping" and rec.status == 200 and rec.method == "GET"
    rendered = JsonFormatter().format(rec)
    for secret in ("SECRET123", "999999", "abc.def.ghi"):
        assert secret not in rendered


@pytest.mark.parametrize(
    "raw",
    [
        "login password=hunter2 ok",
        'body {"password": "hunter2"}',
        "pin_code=482916",
        "Authorization: Bearer eyJhbGciOi.payload.sig",
        "refresh_token=abc.def.ghi",
        "secret: topsecret",
    ],
)
def test_redact_masks_secrets(raw):
    cleaned = redact(raw)
    for leaked in ("hunter2", "482916", "eyJhbGciOi", "abc.def.ghi", "topsecret"):
        assert leaked not in cleaned


def test_json_formatter_redacts_message_and_extras():
    rec = logging.LogRecord("t", logging.INFO, __file__, 1, "token=%s", ("abc123",), None)
    rec.detail = "pin=123456"
    out = JsonFormatter().format(rec)
    assert "abc123" not in out and "123456" not in out


# -- CSV / Excel formula injection -------------------------------------------


@pytest.mark.parametrize("trigger", ["=", "+", "-", "@"])
def test_formula_prefixes_are_neutralised(trigger):
    assert sanitize_cell(f"{trigger}HYPERLINK(\"http://x\")") == f"'{trigger}HYPERLINK(\"http://x\")"


def test_tab_and_cr_prefixes_neutralised():
    assert sanitize_cell("\t=1+1").startswith("'")
    assert sanitize_cell("\r=1+1").startswith("'")


@pytest.mark.parametrize("safe", ["Ali Traders", "", "12.50", "-5.00", "+92", "0300", None, 7])
def test_normal_values_untouched(safe):
    assert sanitize_cell(safe) == safe


def test_csv_output_cannot_start_with_formula():
    from app.services.export_service import _safe_row

    buf = io.StringIO()
    csv.writer(buf).writerow(_safe_row(["INV-1", "=cmd|' /C calc'!A0", "@SUM(1)", "-5.00"]))
    cells = next(csv.reader(io.StringIO(buf.getvalue())))
    assert cells == ["INV-1", "'=cmd|' /C calc'!A0", "'@SUM(1)", "-5.00"]


def test_excel_cells_are_text_not_formulas():
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=sanitize_cell("=1+1"))
    cell = ws.cell(row=1, column=1)
    assert cell.data_type != "f"
    assert cell.value == "'=1+1"
